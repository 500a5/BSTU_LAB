from math import factorial
from itertools import *
from lab1 import *
import docx

doc = docx.Document()


def readFunck():
    print("Input Z: ")
    return [float(x) for x in input().split()]


def getFirstTable(matr_cof, z):
    basic, stringBasic = getBasicList(matr_cof)
    basic = [basic[i] for i in range(len(basic) - len(matr_cof), len(basic))]
    stringBasic = [stringBasic[i] for i in range(len(stringBasic) - len(matr_cof), len(stringBasic))]
    table = [[0 for y in range(len(matr_cof[0]) + 1)] for x in range(len(basic) + 2)]
    table[0][0] = "Б.п."
    table[0][1] = "С.п."
    for i in range(2, len(table[0])):
        table[0][i] = 'X' + str(i - 1)

    for x in range(len(basic)):
        table[stringBasic[x] + 1][0] = 'X' + str(basic[x] + 1)
    for i in range(1, len(table) - 1):
        table[i][1] = matr_cof[i - 1][len(matr_cof[0]) - 1]
    for i in range(len(matr_cof)):
        for j in range(len(matr_cof[0]) - 1):
            table[i + 1][j + 2] = matr_cof[i][j]
    table[len(table) - 1][0] = 'Z'
    table[len(table) - 1][1] = z[0]
    for i in range(1, len(z)):
        table[len(table) - 1][i + 1] = -z[i]
    return table


def checkSolution(table):
    for i in range(2, len(table[0])):
        if(table[len(table) - 1][i] < 0):
            return False
    return True


def getOpSt(table):
    max = 2
    for i in range(2, len(table[0])):
        if (table[len(table) - 1][i] < 0):
            if(table[len(table) - 1][i] <= table[len(table) - 1][max]):
                max = i
    return max


def nextStep(table):
    indexMin = getOpSt(table)
    indexChange = 0
    min = 9999999
    for i in range(1, len(table) - 1):
        if (table[i][indexMin] > 0):
            if(table[i][1]/table[i][indexMin] < min):
                min = table[i][1]/table[i][indexMin]
                indexChange = i
    if indexChange == 0:
        print("Функция не ограничена")
        exit(5)
    delit = table[indexChange][indexMin]
    for i in range(1, len(table[indexChange])):
        table[indexChange][i] /= delit
    newTable = [[y for y in x] for x in table]
    newTable[indexChange][0] = 'X' + str(indexMin - 1)
    for i in range(1, len(table)):
        if i != indexChange:
            cof = -table[i][indexMin]
            for j in range(1, len(table[i])):
                newTable[i][j] = table[i][j] + table[indexChange][j] * cof
    return newTable


def printTable(table):
    print('='*120)
    doc_table = []
    for s in table:
        row = []
        for el in s:
            if type(el) is float:
                el = str(round(el, 3))
            print("%-6s" % el, end=" ")
            row.append(el)
        print()
        doc_table.append(row)
    print('='*120)

    cols = doc_table[0]
    doc_table = doc_table[1:]

    table = doc.add_table(len(doc_table) + 1, len(doc_table[0]), style='TableGrid')
    for j in range(len(cols)):
        table.cell(0, j).text = cols[j]
    # заполняем таблицу данными
    for i, row in enumerate(doc_table):
        for j, text in enumerate(row):
            # получаем ячейку таблицы
            cell = table.cell(i+1, j)
            # записываем в ячейку данные
            cell.text = str(text)
    doc.save('lab.docx')


def print_solution(table):
    print("Решение")
    solution = [0 for x in range(len(table[0]) - 2)]

    for i in range(1, len(table) - 1):
        print(int(table[i][0][1]))
        solution[int(table[i][0][1]) - 1 ] = table[i][1]
    print(solution)
    print("Значение в  точке максимума")
    print("Zmax =", table[len(table) - 1][1])


if __name__ == "__main__":
    matr_cof = readMatr()
    print("Введите какие переменные базисные или вы хотите сделать их базисными")
    matr_cof = getBasicMatr(matr_cof, [int(x) - 1 for x in input().split()])
    z = readFunck()
    table = getFirstTable(matr_cof, z)
    printTable(table)
    while not checkSolution(table):
        table = nextStep(table)
        printTable(table)
    print_solution(table)
